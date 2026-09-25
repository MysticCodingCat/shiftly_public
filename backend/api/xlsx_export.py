"""班表 Excel 匯出：格式化 .xlsx（凍結窗格、班別色塊、覆蓋列、統計工作表）。"""
from __future__ import annotations

import io
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

WEEKDAY_ZH = ["一", "二", "三", "四", "五", "六", "日"]  # date.weekday()

WORK_PALETTE = ["E3EFEB", "FBF0DD", "E7E7F9", "FDE8EF", "E2F1F8", "EEF4DD", "F9E5E2", "E5F3F0"]
WORK_FONTS = ["0C5F4E", "8A5905", "4740A8", "A83368", "0B6285", "55700E", "A13C2E", "0F766E"]
REST_FILL, REST_FONT = "F1F3F6", "7D8A9C"
LEAVE_FILL, LEAVE_FONT = "F3EAF8", "7B3F9E"

THIN = Side(style="thin", color="DDE3EA")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")


def _shift_styles(shifts: list[dict]) -> dict[str, tuple[str, str]]:
    styles = {}
    wi = 0
    for s in shifts:
        if s.get("kind") == "work":
            styles[s["id"]] = (WORK_PALETTE[wi % 8], WORK_FONTS[wi % 8])
            wi += 1
        elif s.get("kind") == "leave":
            styles[s["id"]] = (LEAVE_FILL, LEAVE_FONT)
        else:
            styles[s["id"]] = (REST_FILL, REST_FONT)
    return styles


def _short(shifts_by_id: dict, sid: str) -> str:
    s = shifts_by_id.get(sid, {})
    if s.get("kind") != "work":
        if "REG" in sid:
            return "例"
        if "FLX" in sid:
            return "休"
        if "NAT" in sid:
            return "國"
        return (s.get("name") or sid)[:2]
    return sid


def _required_for(problem: dict, shift_id: str, d: date) -> int | None:
    holidays = {h["date"] for h in problem.get("calendar", {}).get("holidays", [])}
    ds = d.isoformat()
    wd = ["mon", "tue", "wed", "thu", "fri", "sat", "sun"][d.weekday()]
    req = None
    for dem in problem.get("demands", []):
        if dem.get("shiftId") != shift_id or dem.get("min") is None or dem.get("minQualified") is not None:
            continue
        exc = dem.get("except") or {}
        if exc.get("holidays") and ds in holidays:
            continue
        delta = sum(a.get("delta", 0) for a in dem.get("adjustments", [])
                    if wd in (a.get("daysOfWeek") or []) or ds in (a.get("dates") or []))
        need = max(0, dem["min"] + delta)
        req = need if req is None else max(req, need)
    return req


def build_xlsx(problem: dict, result: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "班表"

    shifts = problem.get("shifts", [])
    shifts_by_id = {s["id"]: s for s in shifts}
    styles = _shift_styles(shifts)
    employees = problem.get("employees", [])
    holidays = {h["date"] for h in problem.get("calendar", {}).get("holidays", [])}

    assignments = result.get("assignments", [])
    dates = sorted({a["date"] for a in assignments})
    by_emp: dict[str, dict[str, str]] = {}
    for a in assignments:
        by_emp.setdefault(a["employeeId"], {})[a["date"]] = a["shiftId"]

    # 標題
    title = problem.get("meta", {}).get("name") or "班表"
    ws.cell(row=1, column=1, value=title).font = Font(size=14, bold=True)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=min(8, len(dates) + 2))

    # 表頭
    head_fill = PatternFill("solid", fgColor="FAFBFC")
    weekend_fill = PatternFill("solid", fgColor="F3F1EC")
    ws.cell(row=3, column=1, value="編號")
    ws.cell(row=3, column=2, value="姓名")
    for j, ds in enumerate(dates, start=3):
        d = date.fromisoformat(ds)
        c = ws.cell(row=3, column=j, value=f"{d.day}\n{WEEKDAY_ZH[d.weekday()]}")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.font = Font(size=9, bold=True)
        c.fill = weekend_fill if (d.weekday() >= 5 or ds in holidays) else head_fill
        c.border = BORDER
        ws.column_dimensions[get_column_letter(j)].width = 4.5
    for col, w in ((1, 9), (2, 12)):
        ws.column_dimensions[get_column_letter(col)].width = w
        ws.cell(row=3, column=col).font = Font(size=10, bold=True)
        ws.cell(row=3, column=col).border = BORDER
    ws.row_dimensions[3].height = 26

    # 人員列
    row = 4
    for e in employees:
        ws.cell(row=row, column=1, value=e["id"]).border = BORDER
        ws.cell(row=row, column=2, value=e.get("name", "")).border = BORDER
        for j, ds in enumerate(dates, start=3):
            sid = by_emp.get(e["id"], {}).get(ds, "")
            c = ws.cell(row=row, column=j, value=_short(shifts_by_id, sid) if sid else "")
            c.alignment = CENTER
            c.border = BORDER
            if sid in styles:
                fill, fg = styles[sid]
                c.fill = PatternFill("solid", fgColor=fill)
                c.font = Font(size=9, bold=True, color=fg)
        row += 1

    # 覆蓋列
    for s in shifts:
        if s.get("kind") != "work":
            continue
        ws.cell(row=row, column=1, value=f"{s['id']} 班人數").font = Font(size=9, color="66788C")
        ws.cell(row=row, column=1).border = BORDER
        ws.cell(row=row, column=2).border = BORDER
        for j, ds in enumerate(dates, start=3):
            d = date.fromisoformat(ds)
            count = sum(1 for e in employees if by_emp.get(e["id"], {}).get(ds) == s["id"])
            req = _required_for(problem, s["id"], d)
            c = ws.cell(row=row, column=j, value=f"{count}/{req}" if req is not None else count)
            c.alignment = CENTER
            c.border = BORDER
            c.font = (Font(size=8, bold=True, color="C0392B")
                      if req is not None and count < req else Font(size=8, color="66788C"))
        row += 1

    ws.freeze_panes = "C4"

    # 統計工作表
    st = wb.create_sheet("統計")
    headers = ["人員", "姓名", "工時", "加班", "休假天數", "週末休"] + [s["id"] for s in shifts]
    for j, h in enumerate(headers, start=1):
        st.cell(row=1, column=j, value=h).font = Font(bold=True, size=10)
    name_by_id = {e["id"]: e.get("name", "") for e in employees}
    for i, stat in enumerate(result.get("stats", []), start=2):
        st.cell(row=i, column=1, value=stat["employeeId"])
        st.cell(row=i, column=2, value=name_by_id.get(stat["employeeId"], ""))
        st.cell(row=i, column=3, value=stat.get("hours", 0))
        st.cell(row=i, column=4, value=stat.get("overtimeHours", 0))
        st.cell(row=i, column=5, value=stat.get("restDays", 0))
        st.cell(row=i, column=6, value=stat.get("weekendRestDays", 0))
        for j, s in enumerate(shifts, start=7):
            st.cell(row=i, column=j, value=stat.get("shiftCounts", {}).get(s["id"], 0))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
